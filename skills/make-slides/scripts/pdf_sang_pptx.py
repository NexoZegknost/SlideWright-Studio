#!/usr/bin/env python3
"""pdf_sang_pptx.py — Chuyển PDF (bản trình chiếu xuất ra PDF) thành PowerPoint GIỮ NGUYÊN BẢN GỐC.

Mỗi trang PDF -> 1 slide:
  * Ảnh nền = trang PDF đã xoá chữ và xoá ảnh rời (giữ hình nền, khung, bảng, biểu đồ vẽ bằng nét).
  * Chữ dựng lại thành hộp văn bản sửa được: đúng vị trí, phông, cỡ, màu, đậm/nghiêng, khoảng dòng.
  * Ảnh rời đặt lại đúng vị trí; ảnh có lớp trong suốt (SMask, kể cả /Matte) giữ được nền trong suốt.
  * Hiệu ứng xuất hiện (Fade) theo thứ tự đọc; tiêu đề, phụ đề, số trang hiện sẵn.
    Trang bìa / trang ít nội dung: tự chạy nối tiếp; trang nội dung: mỗi cú nhấp một khối.
  * Chuyển trang Fade.

Cài: pip install --break-system-packages pymupdf python-pptx numpy pillow
Dùng: python3 pdf_sang_pptx.py vao.pdf ra.pptx [--zoom 2.5] [--no-anim]
Sau khi chạy: render lại bằng soffice rồi so ảnh với PDF gốc (xem references/chuyen-doi-pdf.md).
"""
import argparse, io, re
import numpy as np
import pymupdf as fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from lxml import etree

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
GAP_MERGE = 13  # pt: các khối sát nhau theo chiều dọc gộp thành 1 cú nhấp


def font_family(raw):
    name = raw.split("+", 1)[-1]
    fam = re.split(r"[-,]", name)[0]
    fam = re.sub(r"(MT|PS|PSMT)$", "", fam)
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", fam).strip() or "Arial"


def rgba_with_mask(doc, img_bytes, smask_xref):
    rgb = np.asarray(Image.open(io.BytesIO(img_bytes)).convert("RGB")).astype(float)
    m = fitz.Pixmap(doc, smask_xref)
    a = np.frombuffer(m.samples, dtype=np.uint8).reshape(m.height, m.width).astype(float)
    if a.shape != rgb.shape[:2]:
        a = np.asarray(Image.fromarray(a.astype(np.uint8)).resize((rgb.shape[1], rgb.shape[0]))).astype(float)
    if "/Matte" in doc.xref_object(smask_xref):  # bỏ premultiply trên nền đen
        rgb = np.clip(rgb / (np.maximum(a, 1)[..., None] / 255.0), 0, 255)
    out = io.BytesIO()
    Image.fromarray(np.dstack([rgb, a]).astype(np.uint8), "RGBA").save(out, "PNG")
    return out.getvalue()


def render_background(src, pno, remove_xrefs, zoom):
    d = fitz.open(src)
    pg = d[pno]
    for b in pg.get_text("rawdict")["blocks"]:
        if b["type"] == 0:
            for l in b["lines"]:
                for s in l["spans"]:
                    pg.add_redact_annot(fitz.Rect(s["bbox"]), fill=False)
    pg.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                        text=fitz.PDF_REDACT_TEXT_REMOVE)
    for x in remove_xrefs:
        try:
            pg.delete_image(x)
        except Exception:
            pass
    pix = pg.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    return io.BytesIO(pix.tobytes("jpg", jpg_quality=92))


def add_textbox(slide, lines):
    s0 = lines[0]["spans"][0]
    size = s0["size"]
    x0 = min(l["bbox"][0] for l in lines)
    x1 = max(l["bbox"][2] for l in lines)
    top = s0["origin"][1] - 1.0 * size
    pitch = (lines[-1]["spans"][0]["origin"][1] - s0["origin"][1]) / (len(lines) - 1) if len(lines) > 1 else None
    bottom = lines[-1]["spans"][0]["origin"][1] + 0.3 * size
    tb = slide.shapes.add_textbox(Pt(x0), Pt(top), Pt(x1 - x0 + size * 0.6), Pt(bottom - top))
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size = False, None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, l in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if pitch:
            para.line_spacing = Pt(pitch)
        spans = [s for s in l["spans"] if s["text"]]
        for j, s in enumerate(spans):
            r = para.add_run()
            r.text = s["text"].rstrip() if j == len(spans) - 1 else s["text"]
            f = r.font
            fam = font_family(s["font"])
            f.name = fam
            f.size = Pt(round(s["size"] * 2) / 2)
            f.bold = "Bold" in s["font"] or bool(s["flags"] & 16)
            f.italic = "Italic" in s["font"] or "Oblique" in s["font"] or bool(s["flags"] & 2)
            c = s["color"]
            f.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
            rPr = r._r.get_or_add_rPr()
            rPr.set("lang", "vi-VN")
            for tag in ("latin", "cs"):
                el = rPr.find(f"{{{A}}}{tag}")
                if el is None:
                    el = etree.SubElement(rPr, f"{{{A}}}{tag}")
                el.set("typeface", fam)
    return tb


def block_textboxes(slide, b):
    stacks = []
    for l in b["lines"]:
        if not "".join(s["text"] for s in l["spans"]).strip():
            continue
        for st in stacks:
            last = st[-1]
            if (abs(last["bbox"][0] - l["bbox"][0]) < 2
                    and 0 < l["spans"][0]["origin"][1] - last["spans"][0]["origin"][1] < 2 * l["spans"][0]["size"]
                    and abs(last["spans"][0]["size"] - l["spans"][0]["size"]) < 0.3):
                st.append(l)
                break
        else:
            stacks.append([l])
    return [add_textbox(slide, st) for st in stacks]


def timing_xml(groups, auto):
    n = [2]

    def c():
        n[0] += 1
        return n[0]

    def eff(sp, node, delay=0):
        a, b = c(), c()
        return (f'<p:par><p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}">'
                f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c()}" dur="500"/><p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                f'</p:childTnLst></p:cTn></p:par>')

    outer = []
    if auto:
        inner = [f'<p:par><p:cTn id="{c()}" fill="hold"><p:stCondLst><p:cond delay="{gi * 600}"/></p:stCondLst><p:childTnLst>'
                 + "".join(eff(sp, "afterEffect" if k == 0 else "withEffect") for k, sp in enumerate(g))
                 + "</p:childTnLst></p:cTn></p:par>" for gi, g in enumerate(groups)]
        outer.append(f'<p:par><p:cTn id="{c()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst>'
                     f'<p:childTnLst>{"".join(inner)}</p:childTnLst></p:cTn></p:par>')
    else:
        for g in groups:
            effs = "".join(eff(sp, "clickEffect" if k == 0 else "withEffect") for k, sp in enumerate(g))
            outer.append(f'<p:par><p:cTn id="{c()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                         f'<p:par><p:cTn id="{c()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{effs}</p:childTnLst></p:cTn></p:par>'
                         f'</p:childTnLst></p:cTn></p:par>')
    return (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(outer)}</p:childTnLst></p:cTn>'
            f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
            f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src"); ap.add_argument("out")
    ap.add_argument("--zoom", type=float, default=2.5)
    ap.add_argument("--no-anim", action="store_true")
    a = ap.parse_args()
    doc = fitz.open(a.src)
    W, H = doc[0].rect.width, doc[0].rect.height
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    fonts = set()
    for pno, page in enumerate(doc):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        infos = page.get_image_info(xrefs=True)
        sep, bg_full = [], False
        for inf in infos:
            r = fitz.Rect(inf["bbox"]) & page.rect
            if r.width * r.height > 0.85 * W * H:
                bg_full = True
            elif inf["xref"]:
                sep.append(inf)
        slide.shapes.add_picture(render_background(a.src, pno, [i["xref"] for i in sep], a.zoom), 0, 0,
                                 prs.slide_width, prs.slide_height).name = "Hình nền"
        items = []
        for b in page.get_text("dict")["blocks"]:
            if b["type"] == 0:
                items.append(("text", fitz.Rect(b["bbox"]), b))
            else:
                for inf in sep:
                    if not inf.get("_u") and fitz.Rect(inf["bbox"]).intersects(fitz.Rect(b["bbox"])):
                        inf["_u"] = True; items.append(("image", fitz.Rect(inf["bbox"]), inf)); break
        items += [("image", fitz.Rect(i["bbox"]), i) for i in sep if not i.get("_u")]
        units, title_seen = [], False
        for kind, bbox, pl in items:
            if kind == "image":
                ext = doc.extract_image(pl["xref"])
                data = rgba_with_mask(doc, ext["image"], ext["smask"]) if ext.get("smask") else ext["image"]
                r = fitz.Rect(pl["bbox"])
                shp = slide.shapes.add_picture(io.BytesIO(data), Pt(r.x0), Pt(r.y0), Pt(r.width), Pt(r.height))
                units.append((bbox, [shp.shape_id], None)); continue
            shapes = block_textboxes(slide, pl)
            if not shapes:
                continue
            for l in pl["lines"]:
                for s in l["spans"]:
                    fonts.add(font_family(s["font"]))
            sizes = [s["size"] for l in pl["lines"] for s in l["spans"]]
            txt = "".join(s["text"] for l in pl["lines"] for s in l["spans"]).strip()
            is_num = bbox.y0 > H - 45 and txt.isdigit()
            is_title = not title_seen and max(sizes) >= 26 and bbox.y0 < 0.17 * H
            is_sub = title_seen and bbox.y0 < 0.25 * H and max(sizes) < 16 and not units and not bg_full
            if is_title and not bg_full:
                title_seen = True
            if not (is_num or (is_title and not bg_full) or is_sub):
                fs = pl["lines"][0]["spans"][0]
                units.append((bbox, [s.shape_id for s in shapes], (round(fs["size"], 1), "Bold" in fs["font"], round(pl["lines"][0]["bbox"][0]))))
        groups = []
        for bbox, ids, sig in units:  # đề mục + phần giải thích đi cùng 1 cú nhấp; hàng bảng tách riêng
            if groups:
                pb, pids, psig = groups[-1]
                if -2 < bbox.y0 - pb.y1 < GAP_MERGE and min(pb.x1, bbox.x1) > max(pb.x0, bbox.x0) and (sig is None or psig is None or sig != psig):
                    groups[-1] = (pb | bbox, pids + ids, sig); continue
            groups.append((fitz.Rect(bbox), list(ids), sig))
        sld = slide._element
        sld.append(etree.fromstring(f'<p:transition xmlns:p="{P}" spd="med"><p:fade/></p:transition>'))
        if groups and not a.no_anim:
            sld.append(etree.fromstring(timing_xml([g[1] for g in groups], bg_full or len(items) <= 4)))
    prs.save(a.out)
    print("Đã lưu", a.out, "| số slide:", len(doc), "| phông dùng:", ", ".join(sorted(fonts)))


if __name__ == "__main__":
    main()
