#!/usr/bin/env python3
"""xem_truoc_morph.py — xem trước chuyển trang Morph bằng ảnh động (LibreOffice không chạy được Morph).

Với mỗi slide có Morph, dựng các khung trung gian giống cách PowerPoint nội suy:
  - đối tượng chung "!!tên" (và đối tượng cùng tên ở hai slide) dời vị trí, đổi cỡ, xoay, đổi màu/độ trong theo t;
  - phần còn lại của slide trước mờ dần, của slide sau hiện dần (khung = trộn hai ảnh theo t);
  - ở slide sau, đối tượng có hiệu ứng mở đầu bị ẩn (chúng chỉ hiện sau khi morph xong).
Kết quả: <thu_muc>/morph_<NN>.gif cho từng cặp và <thu_muc>/morph_tong.png (mỗi hàng một cặp, các khung theo thời gian)
để kiểm tra bằng mắt: khối có thật sự trôi/xoay/đổi màu không, có đè chữ không, chuyển động có hài hoà không.

Dùng:  python3 xem_truoc_morph.py Slide.pptx [--ra thu_muc] [--khung 6] [--slide 3,5]
Cần: python-pptx, LibreOffice, pdftoppm, Pillow. Phông của gói nên được cài trước (makeslides.installFontsForPreview).
"""
import copy, os, re, subprocess, sys, tempfile
from lxml import etree
from pptx import Presentation
from PIL import Image, ImageDraw

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}


def morph_slides(prs):
    out = []
    for i, s in enumerate(prs.slides):
        x = etree.tostring(s._element).decode("utf8")
        if "p159:morph" in x or "}morph" in x and "2015/09/main" in x: out.append(i)
    return out


def shapes_by_name(el):
    d = {}
    for sp in el.iter("{%s}sp" % P, "{%s}pic" % P, "{%s}cxnSp" % P):
        nv = sp.find(".//p:cNvPr", NS)
        if nv is not None: d.setdefault(nv.get("name"), sp)
    return d


def entr_ids(el):
    ids = set()
    for c in el.iter("{%s}cTn" % P):
        if c.get("presetClass") == "entr":
            for t in c.iter("{%s}spTgt" % P): ids.add(t.get("spid"))
    return ids


def xf(sp):
    x = sp.find(".//a:xfrm", NS)
    if x is None or x.find("a:off", NS) is None: return None
    o, e = x.find("a:off", NS), x.find("a:ext", NS)
    return [int(o.get("x")), int(o.get("y")), int(e.get("cx")), int(e.get("cy")), int(x.get("rot", "0"))]


def set_xf(sp, v):
    x = sp.find(".//a:xfrm", NS)
    o, e = x.find("a:off", NS), x.find("a:ext", NS)
    o.set("x", str(int(v[0]))); o.set("y", str(int(v[1]))); e.set("cx", str(int(v[2]))); e.set("cy", str(int(v[3])))
    x.set("rot", str(int(v[4])))


def fill(sp):
    c = sp.find("p:spPr/a:solidFill/a:srgbClr", NS)
    if c is None: return None
    a = c.find("a:alpha", NS)
    return c, c.get("val"), int(a.get("val")) if a is not None else 100000


def lerp(a, b, t): return a + (b - a) * t


def lerp_hex(a, b, t):
    pa = [int(a[i:i + 2], 16) for i in (0, 2, 4)]; pb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(lerp(x, y, t)):02X}" for x, y in zip(pa, pb))


def pose(dst_el, src_el_a, src_el_b, names, t):
    """đặt các đối tượng chung của dst_el về tư thế nội suy giữa slide A và B"""
    da, db, dd = shapes_by_name(src_el_a), shapes_by_name(src_el_b), shapes_by_name(dst_el)
    for n in names:
        a, b, d = da[n], db[n], dd.get(n)
        if d is None: continue
        ga, gb = xf(a), xf(b)
        if ga and gb: set_xf(d, [lerp(u, v, t) for u, v in zip(ga, gb)])
        # cỡ chữ cũng được nội suy (Morph co/giãn chữ)
        sa = [r.get("sz") for r in a.iter("{%s}rPr" % A) if r.get("sz")]
        sb = [r.get("sz") for r in b.iter("{%s}rPr" % A) if r.get("sz")]
        if sa and sb:
            for r in d.iter("{%s}rPr" % A):
                if r.get("sz"): r.set("sz", str(int(lerp(int(sa[0]), int(sb[0]), t) / 100) * 100))
        fa, fb, fd = fill(a), fill(b), fill(d)
        if fa and fb and fd:
            fd[0].set("val", lerp_hex(fa[1], fb[1], t))
            al = fd[0].find("a:alpha", NS)
            val = int(lerp(fa[2], fb[2], t))
            if al is None and val < 100000:
                al = etree.SubElement(fd[0], "{%s}alpha" % A)
            if al is not None: al.set("val", str(val))


def strip_anim(el):
    for tag in ("timing", "transition"):
        for n in el.findall("p:%s" % tag, NS): el.remove(n)
    for n in list(el.iter("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")):
        if n.getparent() is el: el.remove(n)


def clone(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    el = new._element
    for ch in list(el): el.remove(ch)
    for ch in src._element: el.append(copy.deepcopy(ch))
    for r in src.part.rels.values():
        if "notesSlide" in r.reltype or "slideLayout" in r.reltype: continue
        new.part.rels._rels[r.rId] = r
    return new


def main():
    a = sys.argv[1:]
    if not a: print(__doc__); sys.exit(2)
    src = a[0]
    out = a[a.index("--ra") + 1] if "--ra" in a else "xem_truoc_morph"
    nk = int(a[a.index("--khung") + 1]) if "--khung" in a else 6
    only = {int(x) for x in a[a.index("--slide") + 1].split(",")} if "--slide" in a else None
    os.makedirs(out, exist_ok=True)
    prs = Presentation(src)
    orig = list(prs.slides)
    pairs = [i for i in morph_slides(prs) if i > 0 and (only is None or i + 1 in only)]
    if not pairs: print("Không có slide nào dùng Morph."); return
    ts = [k / (nk - 1) for k in range(nk)]
    plan = []
    for i in pairs:
        A_, B_ = orig[i - 1], orig[i]
        da, db = shapes_by_name(A_._element), shapes_by_name(B_._element)
        names = sorted(n for n in set(da) & set(db) if n.startswith("!!") or (xf(da[n]) and xf(db[n]) and n not in ("title",)))
        names = [n for n in names if n.startswith("!!")] or names
        hide = entr_ids(B_._element)
        for t in ts:
            for which, base in (("a", A_), ("b", B_)):
                s = clone(prs, base)
                strip_anim(s._element)
                pose(s._element, A_._element, B_._element, names, t)
                if which == "b":
                    for sp in list(s._element.iter("{%s}sp" % P, "{%s}pic" % P, "{%s}graphicFrame" % P)):
                        nv = sp.find(".//p:cNvPr", NS)
                        if nv is not None and nv.get("id") in hide and nv.get("name") not in names: sp.getparent().remove(sp)
                plan.append((i, t, which))
    # chỉ giữ các slide khung
    sld = prs.slides._sldIdLst
    for sid in list(sld)[: len(orig)]: sld.remove(sid)
    tmp = tempfile.mkdtemp()
    f = os.path.join(tmp, "khung.pptx"); prs.save(f)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, f], check=True, capture_output=True)
    subprocess.run(["pdftoppm", "-r", "48", "-png", os.path.join(tmp, "khung.pdf"), os.path.join(tmp, "k")], check=True)
    imgs = sorted([x for x in os.listdir(tmp) if x.startswith("k-") and x.endswith(".png")], key=lambda x: int(re.findall(r"\d+", x)[-1]))
    imgs = [Image.open(os.path.join(tmp, x)).convert("RGB") for x in imgs]
    rows = {}
    for k in range(0, len(plan), 2):
        (i, t, _), ia, ib = plan[k], imgs[k], imgs[k + 1]
        rows.setdefault(i, []).append(Image.blend(ia, ib, t))
    w, h = imgs[0].size
    sheet = Image.new("RGB", (w * nk + 8 * (nk + 1), (h + 26) * len(rows) + 8), (245, 242, 236))
    dr = ImageDraw.Draw(sheet)
    for r, (i, frames) in enumerate(rows.items()):
        y = 8 + r * (h + 26)
        dr.text((8, y), f"Morph slide {i} -> {i + 1}", fill=(60, 40, 40))
        for c, fr in enumerate(frames): sheet.paste(fr, (8 + c * (w + 8), y + 18))
        # ảnh động: dừng ở hai đầu cho dễ nhìn
        seq = [frames[0]] * 3 + frames + [frames[-1]] * 4
        seq[0].save(os.path.join(out, f"morph_{i + 1:02d}.gif"), save_all=True, append_images=seq[1:], duration=160, loop=0)
    sheet.save(os.path.join(out, "morph_tong.png"))
    print(f"{len(rows)} chuyển trang Morph → {out}/morph_tong.png và morph_NN.gif")


if __name__ == "__main__":
    main()
